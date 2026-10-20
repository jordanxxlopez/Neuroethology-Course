#!/usr/bin/env python3
"""Build the NEUR 411 course website from the syllabus and the lecture manifest.

Each semester has its own sources; the default is Fall 2026 in source/, other terms in source/<term>/:
  site.json                term name, year, syllabus file, field-trip label, links to the other semesters
  <syllabus>.docx          the syllabus; every policy, the calendar and the readings come from it
  lectures.json            one Google Drive file per lecture deck (and per scheduled special topic)
  announcements.json       announcements, with their HTML in announcements/
Shared by every term (in source/): special_topics.json, class_resources.html, university_resources.html.

Output: docs/ for Fall 2026, docs/<term>/ for other terms (served by GitHub Pages).
Run:  python3 tools/build_site.py            (Fall 2026)
      python3 tools/build_site.py spring-2027
Requires: python-docx. Optional: LibreOffice (soffice) to refresh the syllabus PDF.
"""
import datetime as dt
import html
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "source"          # shared sources
TERM_SRC = SRC                  # this semester's sources
OUT = ROOT / "docs"
SYLLABUS = SRC / "JL_NEUR411.docx"
YEAR = 2026
TERM = "Fall 2026"
A = ""                          # path prefix to the shared docs/assets folder
STUB = "NEUR411_Syllabus_Fall2026"


def configure(term_dir):
    """Point the build at one semester's sources and output folder; returns its settings."""
    global TERM_SRC, OUT, SYLLABUS, YEAR, TERM, A, STUB
    if term_dir:
        TERM_SRC, OUT = SRC / term_dir, ROOT / "docs" / term_dir
    settings = json.loads((TERM_SRC / "site.json").read_text())
    SYLLABUS = TERM_SRC / settings["syllabus"]
    YEAR, TERM, A = settings["year"], settings["term"], settings.get("asset_prefix", "")
    STUB = "NEUR411_Syllabus_" + TERM.replace(" ", "")
    return settings

# Syllabus sections that get their own page instead of appearing on the syllabus page.
OWN_PAGE = {"Course Calendar", "Readings Calendar"}
FIELD_START = "Required Field Experience"  # paragraph in Course Description where the field-trip text begins
def banner():
    return (f'<figure class="banner"><img src="{A}assets/syllabus-banner.jpg" width="1600" height="532" '
            'alt="Photo strip of eleven animals: grizzly bear, mountain lion, bison, coyote, orca, California condor, '
            'gray wolf, sea otter, black wolf, manatee, and a polar bear with its cub."></figure>')

esc = html.escape


# ---------------------------------------------------------------- docx → HTML

URL_RE = re.compile(r"(https?://[^\s<]+[^\s<.,;:)\]])")
EMAIL_RE = re.compile(r"(?<![\w.@/])([\w.+-]+@[\w-]+(?:\.[\w-]+)+)")


def linkify(text_html):
    """Turn bare URLs and e-mail addresses (already HTML-escaped) into links."""
    text_html = URL_RE.sub(lambda m: f'<a href="{m.group(1)}">{m.group(1)}</a>', text_html)
    return EMAIL_RE.sub(lambda m: f'<a href="mailto:{m.group(1)}">{m.group(1)}</a>', text_html)


def runs_html(par):
    """Inline HTML for a paragraph: bold/italic runs and hyperlinks, URLs linkified."""
    pieces = []  # (html, bold, italic, is_link)
    for item in par.iter_inner_content():
        if isinstance(item, Hyperlink):
            text = item.text
            href = item.address or ""
            if not href and item.fragment:
                href = "#" + item.fragment
            if href.startswith("mailto:") or href.startswith("http"):
                pieces.append((f'<a href="{esc(href)}">{esc(text)}</a>', False, False, True))
            else:
                pieces.append((esc(text), False, False, False))
        else:
            if not item.text:
                continue
            pieces.append((esc(item.text), bool(item.bold), bool(item.italic), False))
    # merge neighbouring runs with the same formatting so tags are not fragmented
    merged = []
    for h, b, i, link in pieces:
        if merged and not link and not merged[-1][3] and merged[-1][1:3] == (b, i):
            merged[-1] = (merged[-1][0] + h, b, i, False)
        else:
            merged.append((h, b, i, link))
    out = []
    for h, b, i, link in merged:
        if not link:
            h = linkify(h)
        if not h.strip():
            out.append(h)
            continue
        if i:
            h = f"<em>{h}</em>"
        if b:
            h = f"<strong>{h}</strong>"
        out.append(h)
    s = "".join(out)
    s = re.sub(r"</strong>(\s*)<strong>", r"\1", s)
    return s.strip()


def is_list(par):
    ppr = par._p.pPr
    return ppr is not None and ppr.numPr is not None


def fully_bold(par):
    runs = [r for r in par.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)


def cell_html(cell):
    parts = [runs_html(p) for p in cell.paragraphs]
    return "<br>".join(p for p in parts if p)


def table_html(table, caption=None, hide_caption=False):
    rows = table.rows
    head = "".join(f'<th scope="col">{cell_html(c)}</th>' for c in rows[0].cells)
    body = []
    for r in rows[1:]:
        cells = [cell_html(c) for c in r.cells]
        if not any(cells):
            continue
        tds = [f'<th scope="row">{cells[0]}</th>'] + [f"<td>{c}</td>" for c in cells[1:]]
        body.append("<tr>" + "".join(tds) + "</tr>")
    cls = ' class="visually-hidden"' if hide_caption else ""
    cap = f"<caption{cls}>{esc(caption)}</caption>" if caption else ""
    return wrap_table(f"<table>{cap}<thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>", caption)


_region = [0]


def wrap_table(table, label=None):
    """A keyboard-focusable scroll region so wide tables stay usable on small screens."""
    _region[0] += 1
    lab = f' aria-label="{esc(label)} (table)"' if label else ' aria-label="Table"'
    return f'<div class="table-wrap" role="region" tabindex="0"{lab}>{table}</div>'


def iter_body(doc):
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            yield Table(child, doc)


def read_sections(doc):
    """Split the syllabus at Heading 2 into [(title, [blocks])]; title None = front matter."""
    sections = [(None, [])]
    for block in iter_body(doc):
        if isinstance(block, Paragraph) and block.style.name == "Heading 2":
            title = block.text.strip()
            if title:
                sections.append((title, []))
            continue
        sections[-1][1].append(block)
    return sections


def blocks_html(blocks, base_level=3, table_captions=None, label=None):
    """Render a section's paragraphs, bullet lists, sub-headings and tables."""
    out, in_list = [], False
    captions = iter(table_captions or [])
    last_heading = None
    used = set()
    for b in blocks:
        if isinstance(b, Table):
            if in_list:
                out.append("</ul>")
                in_list = False
            cap = next(captions, None) or last_heading
            hide = cap is None
            cap = cap or label
            if cap in used:
                cap = None
            used.add(cap)
            out.append(table_html(b, cap, hide_caption=hide))
            continue
        text = b.text.strip()
        if not text:
            continue
        if is_list(b):
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{runs_html(b)}</li>")
            continue
        if in_list:
            out.append("</ul>")
            in_list = False
        if b.style.name == "Heading 3" or (fully_bold(b) and len(text) < 90):
            last_heading = text
            level = base_level
            out.append(f'<h{level} id="{slug(text.replace("(", "").replace(")", ""))}">{esc(text)}</h{level}>')
        else:
            out.append(f"<p>{runs_html(b)}</p>")
    if in_list:
        out.append("</ul>")
    return "\n".join(out)


def slug(text):
    s = re.sub(r"\(.*?\)", "", text.lower())
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s or "section"


# ---------------------------------------------------------------- calendar and readings

MONTHS = {m: i for i, m in enumerate(
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}


def parse_date(label):
    m = re.match(r"(Mon|Tue|Wed|Thu|Fri|Sat|Sun)\s+([A-Z][a-z]{2})\s+(\d{1,2})", label)
    if not m:
        return None
    return dt.date(YEAR, MONTHS[m.group(2)], int(m.group(3)))


def row_kind(session):
    s = session.lower()
    if s.startswith("lecture"):
        return "lecture"
    if s.startswith("exam"):
        return "exam"
    if s.startswith("no class"):
        return "noclass"
    if re.match(r"special topic \d", s):
        return "st"
    if s.startswith("special"):
        return "special"
    if s.startswith("lab") or s.startswith("make"):
        return "lab"
    return "due"


def read_calendar(table):
    rows = []
    for r in table.rows[1:]:
        cells = r.cells
        date_label = cells[0].text.strip()
        session = " ".join(cells[1].text.split())
        rows.append({
            "date_label": date_label,
            "date": parse_date(date_label),
            "session": session,
            "kind": row_kind(session),
            "topic_html": cell_html(cells[2]),
            "topic": " ".join(cells[2].paragraphs[0].text.split()),
            "notes_html": cell_html(cells[3]),
        })
    return rows


def item_key(session):
    """("L", n) for Lecture n, ("ST", n) for Special Topic n."""
    m = re.search(r"(\d+)", session)
    return ("ST" if session.lower().startswith("special") else "L", int(m.group(1))) if m else None


def item_id(key):
    return f"lecture-{key[1]}" if key[0] == "L" else f"special-topic-{key[1]}"


def item_name(key):
    return f"Lecture {key[1]}" if key[0] == "L" else f"Special Topic {key[1]}"


def read_readings(table):
    out = {}
    for r in table.rows[1:]:
        key = item_key(" ".join(r.cells[1].text.split()))
        if not key:
            continue
        refs = []
        for p in r.cells[3].paragraphs:
            t = p.text.strip()
            if t:
                refs.append(t)
        out[key] = refs
    return out


def reading_html(ref):
    """'Authors (year). Title. Journal ... https://doi.org/x' → structured citation."""
    m = re.search(r"\s*(https?://\S+)\s*$", ref)
    url = None
    if m:
        url = m.group(1)
        ref = ref[: m.start()]
    body = esc(ref.strip())
    am = re.match(r"(.+?\(\d{4}[a-z]?\)\.)\s*(.*)", body)
    if am:
        body = f'<span class="authors">{am.group(1)}</span> {am.group(2)}'
    link = ""
    if url:
        label = url.replace("https://doi.org/", "doi:") if "doi.org/" in url else url
        link = f' <a class="doi" href="{esc(url)}">{esc(label)}</a>'
    return f"<li>{body}{link}</li>"


# ---------------------------------------------------------------- page shell

NAV = [
    ("index.html", "Home"),
    ("announcements.html", "Announcements"),
    ("syllabus.html", "Syllabus"),
    ("schedule.html", "Schedule"),
    ("lectures.html", "Lectures"),
    ("readings.html", "Readings"),
    ("labs.html", "Labs"),
    ("project.html", "Lecture Project"),
    ("field-experience.html", "Field Experience"),
    ("special-topics.html", "Special Topics"),
    ("resources.html", "Resources"),
]
RESOURCE_TABS = [
    ("resources.html", "Class Resources"),
    ("resources-university.html", "University Policies &amp; Resources"),
]


def page(filename, title, body, settings, description="", section=None):
    nav = []
    for href, label in NAV:
        cur = ' aria-current="page"' if href == (section or filename) else ""
        nav.append(f'<li><a href="{href}"{cur}>{label}</a></li>')
    robots = '<meta name="robots" content="noindex, nofollow">\n' if not settings["public"] else ""
    full_title = (f"{title} | NEUR 411 Neuroethology, {TERM}" if filename != "index.html"
                  else f"NEUR 411 Neuroethology, {TERM}")
    terms = " ".join(f'<a href="{t["href"]}"{" aria-current=\"page\"" if t["label"] == TERM else ""}>{t["label"]}</a>'
                     for t in settings.get("terms", []))
    switch = f'<p class="site-terms"><span class="meta">Semester:</span> {terms}</p>' if terms else ""
    built = dt.date.today().strftime("%B %-d, %Y")
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="color-scheme" content="light only">
{robots}<title>{full_title}</title>
<meta name="description" content="{esc(description or f'Course website for NEUR 411 Neuroethology, {TERM}.')}">
<link rel="stylesheet" href="{A}assets/site.css">
</head>
<body>
<a class="skip-link" href="#main">Skip to main content</a>
<header class="site-header">
  <div class="wrap">
    <p class="site-title"><a href="index.html">NEUR 411 <span class="sep" aria-hidden="true">·</span> Neuroethology</a></p>
    <p class="site-term">{TERM} <span aria-hidden="true">·</span> Jordan Lopez, M.D., Ph.D.</p>
    {switch}
  </div>
  <nav class="site-nav" aria-label="Main">
    <ul class="wrap">{''.join(nav)}</ul>
  </nav>
</header>
<main id="main" tabindex="-1">
<div class="wrap">
{body}
</div>
</main>
<footer class="site-footer">
  <div class="wrap">
    <p>NEUR 411 Neuroethology, {TERM}. Instructor: Jordan Lopez, M.D., Ph.D.,
    <a href="mailto:jlopez11@usc.edu">jlopez11@usc.edu</a>. Grades and submissions are handled in {settings['lms']}.</p>
    <p><a href="accessibility.html">Accessibility</a> <span aria-hidden="true">·</span> Last updated {built}</p>
  </div>
</footer>
</body>
</html>
"""


def fmt_size(n):
    return f"{n / 1_000_000:.1f} MB"


def section_html(sections, title, level=2, sid=None):
    for t, blocks in sections:
        if t == title:
            return (f'<section aria-labelledby="{sid or slug(t)}"><h{level} id="{sid or slug(t)}">{esc(t)}</h{level}>\n'
                    f"{blocks_html(blocks, base_level=level + 1, label=t)}</section>")
    raise KeyError(title)


# ---------------------------------------------------------------- pages

def load_announcements():
    data = json.loads((TERM_SRC / "announcements.json").read_text())
    items = []
    for a in data["announcements"]:
        d = dt.date.fromisoformat(a["date"])
        body = (TERM_SRC / "announcements" / a["file"]).read_text()
        items.append({**a, "d": d, "when": d.strftime("%A, %B %-d, %Y"), "body": body})
    return sorted(items, key=lambda a: a["d"], reverse=True)


def latest_announcement_html():
    items = load_announcements()
    if not items:
        return ""
    a = items[0]
    first = re.search(r"<p>(.*?)</p>", a["body"], re.S)
    excerpt = f"<p>{first.group(1)}</p>" if first else ""
    return f"""<section aria-labelledby="latest" class="latest">
<h2 id="latest">Latest announcement</h2>
<h3><a href="announcements.html#{a['id']}">{esc(a['title'])}</a></h3>
<p class="meta"><time datetime="{a['d'].isoformat()}">{a['when']}</time></p>
{excerpt}
<p><a href="announcements.html#{a['id']}">Read the full announcement<span class="visually-hidden">: {esc(a['title'])}</span></a>
<span aria-hidden="true">·</span> <a href="announcements.html">All announcements</a></p>
</section>"""


def build_announcements(settings):
    parts = []
    for a in load_announcements():
        parts.append(f"""<article class="announcement" aria-labelledby="{a['id']}">
<h2 id="{a['id']}">{esc(a['title'])}</h2>
<p class="meta"><time datetime="{a['d'].isoformat()}">{a['when']}</time></p>
{a['body']}
</article>""")
    body = f"""
<h1>Announcements</h1>
<p class="lede">Course announcements, newest first. Grades are in {settings['lms']}.</p>
{''.join(parts) or '<p>No announcements yet.</p>'}
"""
    return page("announcements.html", "Announcements", body, settings,
                f"Announcements for NEUR 411 Neuroethology, {TERM}.")


def build_home(sections, cal, lectures, settings):
    by = dict(sections)
    desc = next(b for b in by["Course Description"] if isinstance(b, Paragraph) and b.text.strip())
    instructor = blocks_html(by["Instructor"])
    times = by["Course Times & Locations"]
    events = []
    for r in cal:
        if r["date"] is None or r["kind"] == "noclass":
            continue
        label = r["topic"]
        if r["kind"] == "lecture":
            label = f'{r["session"]}: {label}'
        elif r["kind"] == "exam":
            label = f'{r["session"]}: {label}'
        elif r["kind"] == "st":
            label = f'{r["session"]}: {label}'
        elif r["kind"] == "special":
            label = f"Field experience: {label}"
        events.append({"d": r["date"].isoformat(), "label": label, "kind": r["kind"],
                       "when": r["date_label"]})
    events_json = json.dumps(events)
    body = f"""
<h1>NEUR 411 Neuroethology</h1>
{banner()}
<p class="lede">{runs_html(desc)}</p>

<div class="grid-2">
<section aria-labelledby="instructor">
<h2 id="instructor">Instructor</h2>
{instructor}
<p>Teaching assistants and the exam proctors are listed in the <a href="syllabus.html#teaching-assistants">syllabus</a>. Contact your TA first for lab, report and grading questions.</p>
</section>
<section aria-labelledby="meetings">
<h2 id="meetings">Course times and locations</h2>
{blocks_html(times, label="Course times and locations")}
</section>
</div>

{latest_announcement_html()}

<section aria-labelledby="materials">
<h2 id="materials">Course materials</h2>
<ul class="link-list">
<li><a href="lectures.html">Lecture slides</a>: all 40 lectures, grouped by exam.</li>
<li><a href="schedule.html">Course calendar</a>: lectures, exams, labs, field experience and due dates.</li>
<li><a href="readings.html">Recommended readings</a>: the primary papers behind each lecture, with DOI links.</li>
<li><a href="labs.html">Labs</a>: lab schedule, report due dates, computational assignments and the reflection essay.</li>
<li><a href="project.html">Lecture project</a>: the simulation or APA paper options and milestones.</li>
<li><a href="field-experience.html">Field experience</a>: the required {settings['field_label']} trip, {settings['field_dates']}.</li>
<li><a href="special-topics.html">Special topics</a>: additional lecture topics you can request to replace a scheduled lecture.</li>
<li><a href="resources.html">Resources</a>: software, library and study help, and USC policies and support services.</li>
<li><a href="syllabus.html">Syllabus</a>: grading, exams, policies; also as a <a href="files/{STUB}.pdf">PDF</a> or <a href="files/{STUB}.docx">Word document</a>.</li>
</ul>
<p>Grades and assignment submissions are handled in {settings['lms']}. This site holds the course materials.</p>
</section>

<section aria-labelledby="upcoming" class="upcoming">
<h2 id="upcoming">Coming up</h2>
<ol id="upcoming-list" class="event-list"></ol>
<p id="upcoming-fallback">See the <a href="schedule.html">course calendar</a> for every lecture, exam and due date.</p>
</section>
<script>
(function () {{
  var events = {events_json};
  var today = new Date(); today.setHours(0, 0, 0, 0);
  var list = document.getElementById('upcoming-list');
  var shown = 0;
  for (var i = 0; i < events.length && shown < 6; i++) {{
    var e = events[i], d = new Date(e.d + 'T00:00:00');
    if (d < today) continue;
    var li = document.createElement('li');
    if (e.kind === 'exam') li.className = 'is-exam';
    var t = document.createElement('time'); t.dateTime = e.d; t.textContent = e.when;
    li.appendChild(t); li.appendChild(document.createTextNode(' ' + e.label));
    list.appendChild(li); shown++;
  }}
  if (!shown) list.hidden = true;
}})();
</script>
"""
    return page("index.html", "Home", body, settings)


def build_syllabus(sections, settings):
    parts, toc = [], []
    for title, blocks in sections:
        if title is None:
            continue
        sid = slug(title)
        if title in OWN_PAGE:
            target = "schedule.html" if title == "Course Calendar" else "readings.html"
            toc.append(f'<li><a href="{target}">{esc(title)}</a> (separate page)</li>')
            continue
        toc.append(f'<li><a href="#{sid}">{esc(title)}</a></li>')
        parts.append(f'<section aria-labelledby="{sid}"><h2 id="{sid}">{esc(title)}</h2>\n{blocks_html(blocks, label=title)}</section>')
    body = f"""
<h1>Syllabus</h1>
{banner()}
<p class="lede">NEUR 411 Neuroethology, {TERM}. Download the full syllabus as a
<a href="files/{STUB}.pdf">PDF</a> or a
<a href="files/{STUB}.docx">Word document</a>.</p>
<nav class="toc" aria-labelledby="toc-heading">
<h2 id="toc-heading">On this page</h2>
<ol>{''.join(toc)}<li><a href="resources-university.html">University Policies and Resources</a> (separate page)</li></ol>
</nav>
{''.join(parts)}
"""
    return page("syllabus.html", "Syllabus", body, settings,
                "Syllabus for NEUR 411 Neuroethology: grading, exams, labs and course policies.")


def build_schedule(sections, cal, settings):
    by = dict(sections)
    intro = blocks_html([b for b in by["Course Calendar"] if isinstance(b, Paragraph)])
    intro = intro.replace("listed in the Lab Schedule above",
                          'listed on the <a href="labs.html#lab-schedule-and-due-dates">Labs</a> page')
    rows = []
    for r in cal:
        session = esc(r["session"])
        topic = r["topic_html"]
        if r["kind"] in ("lecture", "st"):
            topic = f'<a href="lectures.html#{item_id(item_key(r["session"]))}">{topic}</a>'
        when = esc(r["date_label"])
        if r["date"]:
            when = f'<time datetime="{r["date"].isoformat()}">{when}</time>'
        notes = f'<td role="cell" class="notes">{r["notes_html"]}</td>' if r["notes_html"] else '<td role="cell" class="notes empty"></td>'
        rows.append(f'<tr role="row" class="row-{r["kind"]}"><th role="rowheader" scope="row">{when}</th>'
                    f'<td role="cell" class="session">{session}</td><td role="cell">{topic}</td>{notes}</tr>')
    table = (f'<table class="schedule" role="table"><caption>Course calendar, {TERM}. Rows marked Exam are exams; '
             'rows that begin with a week number are lab meetings for that week.</caption>'
             f'<thead><tr role="row"><th role="columnheader" scope="col">Date ({YEAR})</th><th role="columnheader" scope="col">Session</th>'
             '<th role="columnheader" scope="col">Topic</th><th role="columnheader" scope="col">Assessment / notes</th></tr></thead>'
             f'<tbody>{"".join(rows)}</tbody></table>')
    exams_text = " ".join(f'{esc(r["session"])} ({esc(r["date_label"])}) {esc(r["topic"][0].lower() + r["topic"][1:])}.'
                          for r in cal if r["kind"] == "exam")
    body = f"""
<h1>Course calendar</h1>
{intro}
<p>Lecture titles link to the slides. {exams_text}</p>
<div class="table-wrap schedule-wrap">{table}</div>
"""
    return page("schedule.html", "Schedule", body, settings,
                f"Course calendar for NEUR 411 Neuroethology, {TERM}.")


def deck_links(entry, name):
    """Download button for a deck stored in docs/files/lectures/, otherwise a View link to Google Drive."""
    hidden = f'<span class="visually-hidden"> for {name}</span>'
    local = entry.get("local") if entry else None
    path = OUT / "files" / "lectures" / local if local else None
    if not path or not path.exists():
        if not entry or not entry.get("drive_id"):
            return '<span class="meta">Slides will be posted here.</span>'
        fid = entry["drive_id"]
        return (f'<a class="btn" href="https://drive.google.com/file/d/{fid}/view">View slides{hidden}</a> '
                f'<span class="meta">PowerPoint, {fmt_size(entry["size"])}</span>')
    href = f"files/lectures/{local}"
    return (f'<a class="btn" href="{esc(href)}" download>Download slides{hidden}</a> '
            f'<span class="meta">PowerPoint, {fmt_size(path.stat().st_size)}</span>')


def lecture_rows(cal):
    out = []
    for r in cal:
        if r["kind"] == "lecture":
            out.append((int(re.search(r"\d+", r["session"]).group()), r))
    return out


def item_rows(cal):
    """Lectures and scheduled special topics in calendar order: [(key, row)]."""
    return [(item_key(r["session"]), r) for r in cal if r["kind"] in ("lecture", "st")]


def exam_of(key, r):
    if key[0] == "L":
        return (key[1] - 1) // 10 + 1
    m = re.search(r"Exam (\d)", r["notes_html"])
    return int(m.group(1)) if m else None


def build_lectures(cal, lectures, readings, settings):
    manifest = {("L", e["n"]): e for e in lectures["lectures"]}
    manifest.update({("ST", e["st"]): e for e in lectures.get("special_topics", [])})
    blocks, block = {1: [], 2: [], 3: [], 4: []}, 1
    for key, r in item_rows(cal):
        if key[0] == "L":
            block = exam_of(key, r)
        blocks[block].append((key, r))  # a special topic sits with the lectures around it
    exam_dates = {int(re.search(r"\d+", r["session"]).group()): r["date_label"]
                  for r in cal if r["kind"] == "exam"}
    out = []
    for exam, items in blocks.items():
        if exam is None:
            continue
        lo, hi = (exam - 1) * 10 + 1, exam * 10
        lis = []
        for key, r in items:
            name, iid = item_name(key), item_id(key)
            links = deck_links(manifest.get(key), name)
            nread = len(readings.get(key, []))
            rlink = (f' <a href="readings.html#{iid}">Recommended readings'
                     f'<span class="visually-hidden"> for {name}</span></a> <span class="meta">({nread})</span>'
                     if nread else "")
            when = f'<time datetime="{r["date"].isoformat()}">{esc(r["date_label"])}</time>' if r["date"] else ""
            cls = "lecture special" if key[0] == "ST" else "lecture"
            if key[0] == "ST":
                when += f' <span aria-hidden="true">·</span> {r["notes_html"]}' 
            lis.append(f"""<li class="{cls}" id="{iid}">
<h3><span class="lnum">{name}</span> {r['topic_html']}</h3>
<p class="meta">{when}</p>
<p class="actions">{links}{rlink}</p>
</li>""")
        out.append(f"""<section aria-labelledby="exam-{exam}">
<h2 id="exam-{exam}">Lectures {lo} to {hi} <span class="meta">(Exam {exam}, {esc(exam_dates.get(exam, ''))})</span></h2>
<ol class="lecture-list">{''.join(lis)}</ol>
</section>""")
    body = f"""
<h1>Lecture slides</h1>
<p class="lede">Slides for every lecture. <strong>View slides</strong> opens the deck in Google Drive.
To save a copy, use the download button (the downward arrow) at the top of the Drive viewer.</p>
<p>If Drive asks you to sign in, use your USC account (your @usc.edu email). Downloaded files include the
speaker notes and open in PowerPoint, Keynote, Google Slides or LibreOffice Impress. Lecture slides are the
primary source of instruction, and exams are based on lecture content, including material presented in
class that is not on the slides. Updated versions of each deck will replace these as they are posted.</p>
<nav class="toc" aria-label="Lecture groups"><ul class="inline">
<li><a href="#exam-1">Lectures 1 to 10</a></li><li><a href="#exam-2">11 to 20</a></li>
<li><a href="#exam-3">21 to 30</a></li><li><a href="#exam-4">31 to 40</a></li></ul></nav>
{''.join(out)}
"""
    return page("lectures.html", "Lectures", body, settings,
                "Lecture slides for all 40 lectures of NEUR 411 Neuroethology.")


def build_readings(sections, cal, readings, settings):
    by = dict(sections)
    intro = blocks_html([b for b in by["Readings Calendar"] if isinstance(b, Paragraph)])
    out, jump = [], []
    for key, r in item_rows(cal):
        refs = readings.get(key, [])
        if not refs:
            continue
        name, iid = item_name(key), item_id(key)
        when = esc(r["date_label"])
        out.append(f"""<section class="reading-group" aria-labelledby="{iid}">
<h2 id="{iid}"><span class="lnum">{name}</span> {r['topic_html']}</h2>
<p class="meta">{when} <span aria-hidden="true">·</span> <a href="lectures.html#{iid}">Slides<span class="visually-hidden"> for {name}</span></a></p>
<ul class="refs">{''.join(reading_html(x) for x in refs)}</ul>
</section>""")
        label = str(key[1]) if key[0] == "L" else f"ST{key[1]}"
        aria = "" if key[0] == "L" else f' aria-label="Special Topic {key[1]}"'
        jump.append(f'<li><a href="#{iid}"{aria}>{label}</a></li>')
    jump = "".join(jump)
    body = f"""
<h1>Recommended readings</h1>
{intro}
<p>Links go to each paper's DOI. Many journals require signing in through the USC Libraries; off campus, use the library proxy.</p>
<nav class="toc" aria-labelledby="jump"><h2 id="jump" class="visually-hidden">Jump to a lecture</h2>
<p class="meta" aria-hidden="true">Jump to lecture:</p><ul class="inline numbers">{jump}</ul></nav>
{''.join(out)}
"""
    return page("readings.html", "Readings", body, settings,
                "Recommended primary readings for each NEUR 411 lecture.")


def build_labs(sections, settings):
    body = f"""
<h1>Labs</h1>
<p class="lede">The laboratory component is half of your course grade. This page has the lab rules, the lab
schedule with report due dates, the computational assignments and the Captivity and Conservation Reflection.
Submit lab observations to <a href="https://www.gradescope.com/">Gradescope</a> before you leave each lab and
reports in {settings['lms']}.</p>
{section_html(sections, "Laboratory Component")}
{section_html(sections, "Late Work")}
"""
    return page("labs.html", "Labs", body, settings,
                "Lab schedule, report due dates and computational assignments for NEUR 411.")


def build_project(sections, settings):
    body = f"""
<h1>Lecture project</h1>
<p class="lede">The lecture project is 20 percent of the lecture grade. Choose a simulation project or a
10-page APA paper on one animal system covered in class; the choice is due with your proposal.</p>
{section_html(sections, "Lecture Project (20% of Lecture Grade)", sid="lecture-project")}
{section_html(sections, "Academic Integrity and Use of AI Tools")}
<p>Late submissions: a 30% deduction applies to any project submitted after 11:59 PM on the due date, and
projects more than 48 hours late are not accepted (see <a href="labs.html#late-work">Late Work</a>).
Papers for Option B are found through the <a href="resources.html#papers">library and search tools</a>
on the Resources page.</p>
"""
    return page("project.html", "Lecture Project", body, settings,
                "Lecture project options, milestones and AI policy for NEUR 411.")


def build_field(sections, cal, settings):
    blocks = dict(sections)["Course Description"]
    start = next(i for i, b in enumerate(blocks)
                 if isinstance(b, Paragraph) and b.text.strip().startswith(FIELD_START))
    field = blocks[start + 1:]
    rows = []
    for r in cal:
        if r["kind"] != "special":
            continue
        when = f'<time datetime="{r["date"].isoformat()}">{esc(r["date_label"])}</time>' if r["date"] else esc(r["date_label"])
        rows.append(f'<tr><th scope="row">{when}</th><td>{r["topic_html"]}</td></tr>')
    table = ('<table><caption>Field experience days and special topics</caption><thead><tr>'
             f'<th scope="col">Date ({YEAR})</th><th scope="col">Site and special topic</th></tr></thead>'
             f'<tbody>{"".join(rows)}</tbody></table>')
    refl_due = next((r["date_label"] for r in cal if "reflection" in r["session"].lower()), "")
    body = f"""
<h1>Field experience: {settings['field_label']}</h1>
{blocks_html(field, base_level=2, label="Field experience")}
<h2 id="days">The four days</h2>
{wrap_table(table, "Field experience days")}
<p>The trip is evaluated through the Captivity and Conservation Reflection, a 5-page essay due {refl_due}
(details on the <a href="labs.html#laboratory-component">Labs</a> page).</p>
"""
    return page("field-experience.html", "Field Experience", body, settings,
                f"The required {settings['field_label']} field experience for NEUR 411.")


def build_special(settings, cal):
    data = json.loads((SRC / "special_topics.json").read_text())
    total = sum(len(g["topics"]) for g in data["groups"])
    groups = []
    for g in data["groups"]:
        items = []
        for t in g["topics"]:
            note = f' <span class="meta">(also a {settings["field_label"]} field experience topic)</span>' if t.get("field") else ""
            items.append(f"<li>{esc(t['title'])}{note}</li>")
        groups.append(f'<h3 id="{slug(g["name"])}">{esc(g["name"])} <span class="meta">({len(g["topics"])})</span></h3>'
                      f'<ul class="topic-list">{"".join(items)}</ul>')
    jump = "".join(f'<li><a href="#{slug(g["name"])}">{esc(g["name"])}</a></li>' for g in data["groups"])
    sched = [(item_key(r["session"]), r) for r in cal if r["kind"] == "st"]
    scheduled_html = ""
    if sched:
        lis = "".join(f'<li><a href="lectures.html#{item_id(k)}">{item_name(k)}</a> ({esc(r["date_label"])}): '
                      f'{r["topic_html"]} <span class="meta">{r["notes_html"]}</span></li>' for k, r in sched)
        scheduled_html = (f'<section aria-labelledby="this-term"><h2 id="this-term">On the schedule this semester</h2>'
                          f'<p>These special topics are part of the {TERM} schedule. Each one is the lecture on an exam day, '
                          f'and none of them is covered on an exam.</p>'
                          f'<ul>{lis}</ul></section>')
    body = f"""
<h1>Special topics</h1>
<p class="lede">Beyond the 40 scheduled lectures, the course has a library of special-topic lectures on other
animals and questions. They can replace a scheduled lecture during the semester, and every one of them is
a reference for your lecture project and for going further.</p>

{scheduled_html}
<section aria-labelledby="requesting">
<h2 id="requesting">Requesting a special topic</h2>
<ul>
<li>Throughout the semester, you may request that a special topic replace a scheduled lecture. A request can
swap a lecture for the same subject in another animal (for example, a lecture on electroreception could be
replaced with a special topic on electroreception and lateral-line sensing in the axolotl), or for any topic
on the list below.</li>
<li>Special topics are not guaranteed. I decide whether a lecture is replaced, and usually only one to five
lectures in a semester are replaced with special topics.</li>
<li>To make a request, e-mail <a href="mailto:jlopez11@usc.edu">jlopez11@usc.edu</a> with the subject line
“NEUR 411: Special topic request”. Name the topic you want, the scheduled lecture it would replace, and why.</li>
<li>Any replacement is announced in class and posted on the <a href="schedule.html">course calendar</a>
before the lecture.</li>
</ul>
</section>

<section aria-labelledby="topics">
<h2 id="topics">Topics you can choose from <span class="meta">({total} topics)</span></h2>
<nav class="toc" aria-label="Special topic groups"><ul class="inline">{jump}</ul></nav>
{''.join(groups)}
</section>
"""
    return page("special-topics.html", "Special Topics", body, settings,
                "Special-topic lectures students can request to replace a scheduled NEUR 411 lecture.")


def resource_tabs(current):
    items = []
    for href, label in RESOURCE_TABS:
        cur = ' aria-current="page"' if href == current else ""
        items.append(f'<li><a href="{href}"{cur}>{label}</a></li>')
    return f'<nav class="subnav" aria-label="Resources"><ul>{"".join(items)}</ul></nav>'


def build_resources(sections, settings):
    from_syllabus = [section_html(sections, t) for t in
                     ("Teaching Assistants", "Communication", "Course Materials",
                      "Lecture Recordings and Study Materials", "Study Expectations")]
    body = f"""
<h1>Resources</h1>
{resource_tabs("resources.html")}
<p class="lede">Where to get help and the tools you need for this course. University policies and campus
support services are under <a href="resources-university.html">University Policies &amp; Resources</a>.</p>
{(SRC / "class_resources.html").read_text()}
{''.join(from_syllabus)}
"""
    return page("resources.html", "Class Resources", body, settings,
                "Course platforms, software, library and study resources for NEUR 411.", section="resources.html")


def build_university(settings):
    body = f"""
<h1>Resources</h1>
{resource_tabs("resources-university.html")}
<p class="lede">USC policies that apply to every course, and the university's support services for students.
In an emergency, call <a href="tel:+12137404321">(213) 740-4321</a> (UPC) or
<a href="tel:+13234421000">(323) 442-1000</a> (HSC), or 911.</p>
{(SRC / "university_resources.html").read_text()}
"""
    return page("resources-university.html", "University Policies & Resources", body, settings,
                "USC academic integrity, recording, accommodation policies and student support services.",
                section="resources.html")


def build_accessibility(settings):
    body = f"""
<h1>Accessibility</h1>
<p class="lede">This site is built to be usable with a keyboard, a screen reader, browser zoom and high-contrast
display settings.</p>
<h2>How the site is built</h2>
<ul>
<li>Every page has the same navigation, a skip link to the main content and a logical heading structure.</li>
<li>Text and links meet WCAG 2.2 AA contrast; the site always uses a white background, whatever the device's dark-mode setting; links are underlined, and the keyboard focus is always visible.</li>
<li>Tables have captions and header cells. Wide tables scroll inside their own region on small screens.</li>
<li>Pages work without JavaScript. The only script on the site lists upcoming dates on the home page.</li>
<li>The syllabus is available as an HTML page, a PDF and a Word document.</li>
</ul>
<h2>Lecture slides</h2>
<p>Slides open in Google Drive's viewer, which works with screen readers and zoom; the viewer's download
button saves the PowerPoint file, including the speaker notes. Figures in the slides come from published research articles;
each carries its source and citation.</p>
<h2>Requesting another format or reporting a problem</h2>
<p>If any material is hard to use, e-mail <a href="mailto:jlopez11@usc.edu">jlopez11@usc.edu</a> with the page
or lecture number and the format you need. For approved accommodations, contact the Office of Student
Accessibility Services and see the <a href="syllabus.html#accommodations">accommodations policy</a>.</p>
"""
    return page("accessibility.html", "Accessibility", body, settings)


# ---------------------------------------------------------------- main

def refresh_syllabus_files():
    files = OUT / "files"
    files.mkdir(parents=True, exist_ok=True)
    docx_out = files / f"{STUB}.docx"
    shutil.copyfile(SYLLABUS, docx_out)
    pdf_out = files / f"{STUB}.pdf"
    stale = not pdf_out.exists() or pdf_out.stat().st_mtime < SYLLABUS.stat().st_mtime
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if stale and soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(files), str(docx_out)],
                       check=True, capture_output=True)
    elif stale:
        print("warning: syllabus PDF is out of date and LibreOffice is not installed", file=sys.stderr)


def main():
    term_dir = sys.argv[1] if len(sys.argv) > 1 else None
    settings = configure(term_dir)
    lectures = json.loads((TERM_SRC / "lectures.json").read_text())
    doc = Document(SYLLABUS)
    sections = read_sections(doc)
    by = dict(sections)
    cal = read_calendar(next(b for b in by["Course Calendar"] if isinstance(b, Table)))
    readings = read_readings(next(b for b in by["Readings Calendar"] if isinstance(b, Table)))

    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "assignments.html").unlink(missing_ok=True)
    pages = {
        "index.html": build_home(sections, cal, lectures, settings),
        "syllabus.html": build_syllabus(sections, settings),
        "schedule.html": build_schedule(sections, cal, settings),
        "lectures.html": build_lectures(cal, lectures, readings, settings),
        "readings.html": build_readings(sections, cal, readings, settings),
        "labs.html": build_labs(sections, settings),
        "project.html": build_project(sections, settings),
        "field-experience.html": build_field(sections, cal, settings),
        "special-topics.html": build_special(settings, cal),
        "announcements.html": build_announcements(settings),
        "resources.html": build_resources(sections, settings),
        "resources-university.html": build_university(settings),
        "accessibility.html": build_accessibility(settings),
    }
    for name, content in pages.items():
        (OUT / name).write_text(content)
    if not term_dir:  # site-wide files live at the top of docs/
        (OUT / ".nojekyll").write_text("")
        robots = "User-agent: *\nAllow: /\n" if settings["public"] else "User-agent: *\nDisallow: /\n"
        (OUT / "robots.txt").write_text(robots)
    refresh_syllabus_files()

    n_lec = len(lecture_rows(cal))
    n_ref = sum(len(v) for v in readings.values())
    missing = sorted(set(range(1, n_lec + 1)) - {e["n"] for e in lectures["lectures"]})
    n_st = sum(1 for r in cal if r["kind"] == "st")
    print(f"built {TERM} in {OUT.relative_to(ROOT)}: {len(pages)} pages, {n_lec} lectures, {n_st} special topics, {n_ref} readings"
          + (f"; no slides for lectures {missing}" if missing else ""))


if __name__ == "__main__":
    main()
